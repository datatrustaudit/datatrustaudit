import io
import re
from urllib.parse import urlparse

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from datetime import datetime
import openpyxl

from db.database import get_db
from db.models import Scan, User
from api.auth import get_current_active_user

router = APIRouter(prefix="/api/exports", tags=["Exports"])

PAID_PLANS = {"pro", "premium"}


def _campo(item: dict, *nomes, padrao=None):
    """Primeiro nome presente no dicionário.

    raw_data é gravado por jsonable_encoder (snake_case), enquanto a resposta
    da API passa por to_dict() (camelCase). O PDF e o XLSX liam só camelCase e
    por isso mostravam o valor padrão em todas as colunas de tag — "Loaded
    Before Consent?" saía "Não" mesmo num site sem consentimento algum.
    """
    for nome in nomes:
        if nome in item and item[nome] is not None:
            return item[nome]
    return padrao


def _safe_filename_part(url: str) -> str:
    """Sanitiza o host para uso em Content-Disposition.

    O host vem da URL que o usuário mandou auditar: sem filtro, um valor com
    CRLF, aspas ou ponto-e-vírgula quebra (ou injeta) o cabeçalho HTTP.
    """
    host = urlparse(url).netloc or "scan"
    host = re.sub(r"[^A-Za-z0-9._-]", "_", host)
    return host[:60] or "scan"


def _load_paid_scan(scan_id: int, current_user: User, db: Session, formato: str) -> Scan:
    """Valida plano e posse do scan. Compartilhado por XLSX e PDF."""
    if current_user.plan not in PAID_PLANS:
        raise HTTPException(
            status_code=403,
            detail=f"Exportação para {formato} está disponível apenas nos planos Pro e Premium.",
        )

    scan = db.query(Scan).filter(Scan.id == scan_id, Scan.user_id == current_user.id).first()
    if not scan or not scan.raw_data:
        raise HTTPException(status_code=404, detail="Scan data not found or still processing.")
    return scan


# NOTA: as chaves de scan.raw_data são snake_case. raw_data vem de
# jsonable_encoder(audit_result), que serializa a dataclass via asdict e
# ignora o to_dict() camelCase dos models. Ler "estimatedRiskExposure" aqui
# devolve sempre vazio.

@router.get("/xlsx/{scan_id}")
@router.get("/excel/{scan_id}")  # alias: o frontend chama /excel
def export_scan_xlsx(
    scan_id: int,
    current_user: User = Depends(get_current_active_user),
    db: Session = Depends(get_db),
):
    """Exports Scan findings to XLSX. Available only for Pro/Premium."""
    scan = _load_paid_scan(scan_id, current_user, db, "XLSX")

    wb = openpyxl.Workbook()

    # Overview Sheet
    ws1 = wb.active
    ws1.title = "Resumo do Scan"
    ws1.append(["URL Auditada:", scan.url])
    ws1.append(["Data do Scan:", str(scan.completed_at)])
    ws1.append(["Score de Conformidade:", f"{scan.score}/100"])

    data = scan.raw_data or {}
    tags = data.get("tags") or []
    privacy = data.get("privacy") or {}
    duplicates = data.get("duplicates") or []

    ws1.append(["Total de Tags Detectadas:", len(tags)])
    ws1.append(["Ferramenta de Consentimento:", "Sim" if privacy.get("has_consent_tool") else "Não"])
    if privacy.get("total_violations"):
        ws1.append(["Violações de Privacidade:", privacy.get("total_violations")])
        ws1.append(["Exposição Estimada:", privacy.get("estimated_risk_exposure", "")])

    # Tags Sheet
    ws2 = wb.create_sheet("Tags e Analytics")
    ws2.append(["ID", "Name", "Type", "Risk", "Found on Line", "Loaded Before Consent?"])
    for t in tags:
        ws2.append([
            _campo(t, "tagId", "tag_id", padrao="N/A"),
            _campo(t, "name", padrao="Unknown"),
            _campo(t, "type", padrao="Unknown"),
            _campo(t, "lgpdRisk", "privacy_risk", padrao="medium"),
            _campo(t, "lineNumber", "line_number", padrao="Unknown"),
            "Sim" if _campo(t, "isBeforeConsent", "is_before_consent", padrao=False) else "Não"
        ])

    # Privacy Violations Sheet
    if privacy.get("violations"):
        ws3 = wb.create_sheet("Violações de Privacidade")
        ws3.append(["Severidade", "Resumo", "Descrição", "Recomendação"])
        for v in privacy.get("violations"):
            ws3.append([
                v.get("severity", "medium"),
                v.get("summary", ""),
                v.get("description", ""),
                v.get("recommendation", "")
            ])

    # Duplicates Sheet
    if duplicates:
        ws4 = wb.create_sheet("Duplicidades Identificadas")
        ws4.append(["Severidade", "Entidade", "Nome", "Contagem", "Recomendação"])
        for d in duplicates:
            ws4.append([
                d.get("severity", ""),
                d.get("entity_type", ""),
                d.get("name", ""),
                d.get("occurrence_count", 0),
                d.get("recommendation", "")
            ])

    # Cookies e armazenamento gravados sem aceite — a evidência mais direta do
    # laudo. O scanner nunca clica em "aceitar": o que está aqui foi gravado
    # antes de qualquer manifestação do usuário.
    storage = data.get("storage") or {}
    cookies_sem_consentimento = storage.get("cookiesWithoutConsent") or storage.get("cookies_sem_consentimento") or []
    if cookies_sem_consentimento:
        ws5 = wb.create_sheet("Cookies sem Consentimento")
        ws5.append(["Severidade", "Cookie", "Responsável", "Finalidade", "Duração (dias)", "HttpOnly", "Domínio"])
        for c in cookies_sem_consentimento:
            ws5.append([
                c.get("severity") or c.get("severidade", ""),
                c.get("name") or c.get("nome", ""),
                c.get("owner") or c.get("dono", ""),
                c.get("purpose") or c.get("finalidade", ""),
                c.get("durationDays") if c.get("durationDays") is not None else c.get("duracao_dias", ""),
                "Sim" if (c.get("httpOnly") or c.get("http_only")) else "Não",
                c.get("domain") or c.get("dominio", ""),
            ])

    storage_sem_consentimento = storage.get("storageWithoutConsent") or storage.get("storage_sem_consentimento") or []
    if storage_sem_consentimento:
        ws6 = wb.create_sheet("Armazenamento Local")
        ws6.append(["Severidade", "Tipo", "Chave", "Responsável", "Observação"])
        for s_item in storage_sem_consentimento:
            ws6.append([
                s_item.get("severity") or s_item.get("severidade", ""),
                s_item.get("type") or s_item.get("tipo", ""),
                s_item.get("name") or s_item.get("nome", ""),
                s_item.get("owner") or s_item.get("dono", ""),
                s_item.get("detail") or s_item.get("detalhe", ""),
            ])

    # Ordem real de disparo — quando cada rastreador enviou dados, medido pelo
    # relógio da própria visita. É o que separa evidência de indício.
    linha_do_tempo = storage.get("firingTimeline") or storage.get("linha_do_tempo") or []
    if linha_do_tempo:
        ws7 = wb.create_sheet("Ordem de Disparo")
        ws7.append(["#", "Milissegundos", "Responsável", "URL"])
        for e in linha_do_tempo:
            ws7.append([
                e.get("order") or e.get("ordem", ""),
                e.get("msSincePageStart") if e.get("msSincePageStart") is not None else e.get("ms_desde_o_inicio", ""),
                e.get("owner") or e.get("dono", ""),
                e.get("url", ""),
            ])

    # Save to memory stream
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"GTM_Audit_{_safe_filename_part(scan.url)}_{datetime.now().strftime('%Y%m%d')}.xlsx"

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@router.get("/pdf/{scan_id}")
def export_scan_pdf(
    scan_id: int,
    current_user: User = Depends(get_current_active_user),
    db: Session = Depends(get_db),
):
    """Exports Scan findings to PDF. Available only for Pro/Premium."""
    scan = _load_paid_scan(scan_id, current_user, db, "PDF")

    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    stream = io.BytesIO()
    c = canvas.Canvas(stream, pagesize=letter)

    # Title
    c.setFont("Helvetica-Bold", 20)
    c.drawString(50, 750, "DataTrust Audit - Relatório de Scan")

    c.setFont("Helvetica", 12)
    c.drawString(50, 720, f"URL: {scan.url}")
    c.drawString(50, 700, f"Data: {scan.completed_at}")
    c.drawString(50, 680, f"Score de Conformidade: {scan.score}/100")

    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, 640, "Resumo Executivo")
    c.setFont("Helvetica", 12)

    data = scan.raw_data or {}
    privacy = data.get("privacy") or {}
    tags = data.get("tags") or []
    duplicates = data.get("duplicates") or []

    c.drawString(50, 620, f"Total de Tags: {len(tags)}")
    has_cmp = "Sim" if privacy.get("has_consent_tool") else "Não"
    c.drawString(50, 600, f"Gestão de Consentimento Ativa (CMP): {has_cmp}")
    c.drawString(50, 580, f"Violações Encontradas: {privacy.get('total_violations', 0)}")

    if privacy.get("estimated_risk_exposure"):
        c.drawString(50, 560, f"Risco de Exposição: {privacy.get('estimated_risk_exposure')}")

    # Cookies gravados sem aceite: é o achado que sustenta o laudo.
    storage_pdf = data.get("storage") or {}
    cookies_pdf = storage_pdf.get("cookiesWithoutConsent") or storage_pdf.get("cookies_sem_consentimento") or []
    storage_itens_pdf = storage_pdf.get("storageWithoutConsent") or storage_pdf.get("storage_sem_consentimento") or []
    if cookies_pdf or storage_itens_pdf:
        c.drawString(50, 540, f"Cookies gravados sem consentimento: {len(cookies_pdf)}")
        c.drawString(50, 522, f"Chaves de armazenamento local sem consentimento: {len(storage_itens_pdf)}")
        y_tags = 480
    else:
        y_tags = 520

    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, y_tags, "Tags Detectadas")
    c.setFont("Helvetica", 10)

    y = y_tags - 20
    for t in tags[:10]:  # Limit to top 10 for basic PDF to avoid pagination complexity
        risco = _campo(t, 'lgpdRisk', 'privacy_risk', padrao='?')
        antes = 'Sim' if _campo(t, 'isBeforeConsent', 'is_before_consent', padrao=False) else 'Não'
        tag_str = f"- {t.get('name')} ({t.get('type')}) - Risco LGPD: {risco} - Antes do Cookie: {antes}"
        c.drawString(50, y, tag_str)
        y -= 20
        if y < 50:
            break

    if len(tags) > 10:
        c.drawString(50, y, f"... mais {len(tags) - 10} tags encontradas.")

    y -= 40
    if y > 100 and duplicates:
        c.setFont("Helvetica-Bold", 14)
        c.drawString(50, y, "Duplicidades Recorrentes")
        y -= 20
        c.setFont("Helvetica", 10)
        for d in duplicates[:5]:
            vezes = _campo(d, 'occurrenceCount', 'occurrence_count', padrao='?')
            c.drawString(50, y, f"- {d.get('name')} (Risco: {d.get('severity')}) - {vezes}x")
            y -= 20

    c.save()
    stream.seek(0)

    filename = f"GTM_Audit_{_safe_filename_part(scan.url)}.pdf"

    return StreamingResponse(
        stream,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
